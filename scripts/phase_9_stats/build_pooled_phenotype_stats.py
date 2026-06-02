"""Driver — consolidated phenotype-pooled stats suite (kickoff doc, 2026-06-01).

Runs the scalar stats engine (`stats.pooled_endpoint_suite`) over EVERY
phenotype-pooled scalar family, on the pre-stroke-baseline-normalized Δ at
session grain, split by trial cue side, for the 3 post-vs-pre contrasts
(acute/sub_acute/chronic vs pre) and BOTH scopes (3_tier, 4_tier). BH-FDR within
family. Writes ONE consolidated sheet per scope:

    stats/results/phenotype_pooled/<scope>/pooled_phenotype_stats.csv

Families + per-session sources:
  - scalar    : stats/inputs/<ep>/<animal>.parquet (tongue/jaw + spout-behavior
                d_prime / first_lick_accuracy_all / ... ) -> per-session mean.
  - phase_aligned_endpoints : per-trial JSON via
                load_session_phase_aligned_endpoints_long -> per-session mean.
  - wasserstein / normalized_wasserstein : the persisted per-session W-2 CSV.

(Trajectories already carry the full suite incl. cluster permutation, which is
N/A for scalars. Per-trial MWU skipped per Priya.)

Usage::
  python scripts/phase_9_stats/build_pooled_phenotype_stats.py
  python scripts/phase_9_stats/build_pooled_phenotype_stats.py --families scalar
  python scripts/phase_9_stats/build_pooled_phenotype_stats.py --n-boot 5000
"""
from __future__ import annotations

import argparse
import glob
import logging
import sys
from pathlib import Path

import numpy as np
import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from stroke_orofacial.animals import load_animals  # noqa: E402
from stroke_orofacial.cohort_filters import (  # noqa: E402
    DlcSessionFilter,
    build_dlc_pertrial_lick_frame,
    compute_dlc_session_lick_stats,
    load_dlc_filter_spec_from_defaults,
    resolve_dlc_trial_filter,
)
from stroke_orofacial.config_loader import load_defaults  # noqa: E402
from stroke_orofacial.paths import PathResolver  # noqa: E402
from stroke_orofacial.stats import pooled_endpoint_suite as SUITE  # noqa: E402
from stroke_orofacial.stats.aggregators import (  # noqa: E402
    load_session_phase_aligned_endpoints_long,
)
from stroke_orofacial.stats.multiple_comparisons import (  # noqa: E402
    apply_correction_to_results_df,
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                    datefmt="%H:%M:%S")
log = logging.getLogger("build_pooled_phenotype_stats")

PHENOTYPE_REMAP_3TIER = {"moderate-severe": "moderate"}
PHASE_ALIGNED_ENDPOINTS = (
    "peak_angle_deg",
    "signed_angle_at_max_abs_phase_0_075_deg",
    "signed_x_at_max_abs_phase_0_075_px",
)
# Behavior d'/criterion pool across cue sides -> can't split by cue side.
CROSS_CUE_ENDPOINTS = frozenset({"d_prime", "criterion_c"})

# Diagnostic scalar endpoints excluded from the pooled-phenotype hypothesis set
# (Priya 2026-06-01): QC frame-count + pass flag + redundant jaw-deflection
# variants (jaw_peak_abs_deflection is retained as the canonical deflection).
# These are additionally pruned ON TOP of the figured-only scope below; the
# combined effect shrinks the BH-FDR family to genuine hypothesis endpoints.
SCALAR_DIAGNOSTIC_DROP = frozenset({
    "jaw_n_frames_over_thresh",
    "jaw_pass_qc",
    "jaw_peak_neg_deflection",
    "jaw_peak_pos_deflection",
    "jaw_peak_selected_deflection",
})

# Families whose headline significance is the session-MWU (+ bootstrap CI), NOT
# the LMM: the per-session scalar/endpoint cells have only 2-4 animals per
# phenotype, so the animal random-effect variance is frequently non-estimable
# (singular -> convergence_failed). The LMM is reported as a supporting column
# where it converges; trajectories (rich per-trial nesting) keep LMM as their
# headline in their own per-family CSVs. (Priya 2026-06-01.)
MWU_HEADLINE_FAMILIES = frozenset({
    "scalar", "phase_aligned_endpoints",
    "phase_aligned_endpoints_all_licks", "wasserstein_recovery",
    "normalized_wasserstein_recovery",
})
DLC_FILTER_METRIC = "dlc_tongue_trajectory_metrics"


def _figured_scalar_endpoints(resolver: PathResolver) -> set[str]:
    """Slugs that have a phenotype-pooled FIGURE dir (the in-scope set).

    The consolidated sheet's scope is "stats for the figures in the
    phenotype_pooled folder" (Priya). Scalar endpoints with no figure dir
    (the 50 ``spatial_*`` divergence/UOT/wasserstein variants + ``latency_hist``)
    are out of scope and excluded.
    """
    fig = resolver.pipeline_outputs(subdir="stats") / "figures" / "phenotype_pooled"
    if not fig.is_dir():
        return set()
    return {p.name for p in fig.iterdir() if p.is_dir()}


def _session_stats_frame(animal: str, resolver: PathResolver) -> pd.DataFrame:
    """Per-(animal, date, side) DLC v7-kept session lick stats (port of q13b)."""
    from stroke_orofacial.cohort_filters import (
        DLC_SESSION_N_KEPT_LICKS_COL,
        DLC_SESSION_N_TRIALS_TOTAL_COL,
        DLC_SESSION_N_TRIALS_WITH_LICK_COL,
    )
    from stroke_orofacial.cohort_enumeration import yaml_driven_session_dates
    import datetime
    rows: list[dict] = []
    for d in yaml_driven_session_dates(animal):
        if isinstance(d, str):
            d = datetime.date.fromisoformat(d)
        stats = compute_dlc_session_lick_stats(animal, d, resolver)
        if stats is None:
            continue
        for s in stats:
            rows.append({
                "animal": s.animal, "date": s.date, "side": s.side,
                DLC_SESSION_N_TRIALS_TOTAL_COL: s.n_trials_of_side,
                DLC_SESSION_N_TRIALS_WITH_LICK_COL: s.n_trials_with_detected_lick,
                DLC_SESSION_N_KEPT_LICKS_COL: s.n_kept_licks,
            })
    return pd.DataFrame(rows)


def _build_dlc_filter(animals: list[str], resolver: PathResolver):
    """Return ``(trial_filter, kept_keys)`` for the filtered_v1 variant.

    ``kept_keys`` is the set of ``(animal, date, side)`` sessions that pass the
    ``DlcSessionFilter`` (``min_trials_with_detected_lick``); ``None`` if the
    filter is pass-through. ``trial_filter`` gates per-trial rows by
    ``min_licks_per_trial``. This mirrors the on-the-fly filter the v3 figures
    apply in q13b_v2_emit so the stats match the canonical figures.
    """
    defaults = load_defaults()
    trial_filter = resolve_dlc_trial_filter(DLC_FILTER_METRIC, defaults)
    session_filter: DlcSessionFilter = load_dlc_filter_spec_from_defaults(defaults)
    kept_keys: set | None = None
    if not session_filter.is_pass_through():
        sess_frames = [_session_stats_frame(a, resolver) for a in animals]
        sess_frames = [f for f in sess_frames if not f.empty]
        if sess_frames:
            sess_df = pd.concat(sess_frames, ignore_index=True)
            kept_sess, res = session_filter.filter(sess_df)
            kept_keys = {
                (r["animal"], r["date"], r["side"])
                for _, r in kept_sess.iterrows()
            }
            log.info("DLC session filter: %d/%d (animal,date,side) sessions kept",
                     res.n_sessions_kept, res.n_sessions_in)
    return trial_filter, kept_keys


def _apply_trial_then_session_filter(
    long_df: pd.DataFrame, trial_filter, kept_keys: set | None,
) -> pd.DataFrame:
    """Drop low-lick trials then sparse sessions (per-trial long-form in)."""
    out = long_df
    if (trial_filter is not None and not trial_filter.is_pass_through()
            and "n_licks_per_trial" in out.columns):
        keep = (out["n_licks_per_trial"].isna()
                | (out["n_licks_per_trial"] >= trial_filter.min_licks_per_trial))
        out = out[keep]
    if kept_keys is not None:
        out = out[out.apply(
            lambda r: (r["animal"], r["date"], r["side"]) in kept_keys, axis=1)]
    return out.copy()


def _persession_from_pertrial(df: pd.DataFrame, endpoint: str) -> pd.DataFrame:
    """Per-(animal, date, side) mean of finite per-trial values."""
    if df.empty:
        return df
    if "is_sentinel" in df.columns:
        df = df[~df["is_sentinel"].astype(bool)]
    df = df.dropna(subset=["value"])
    g = (df.groupby(["animal", "date", "side"], as_index=False)
           .agg(value=("value", "mean"), epoch_label=("epoch_label", "first")))
    g["endpoint"] = endpoint
    return g


def load_scalar_family(
    resolver: PathResolver, *, filter_variant: str = "filtered_v1",
) -> pd.DataFrame:
    """Scalar endpoints -> per-session long-form, scoped + filtered.

    Scope: only endpoints with a phenotype-pooled FIGURE dir, minus the
    diagnostic drop-set. Input: when ``filter_variant == "filtered_v1"`` each
    endpoint reads ``stats/inputs/filtered_v1/<ep>/`` if that tree exists
    (the QC-filtered tongue family), else falls back to the unfiltered
    ``stats/inputs/<ep>/`` (jaw / spout-behavior have no filtered tree and are
    unfiltered in their canonical figures too).
    """
    inp = resolver.pipeline_outputs(subdir="stats") / "inputs"
    figured = _figured_scalar_endpoints(resolver)
    eps = sorted(
        d.name for d in inp.iterdir()
        if d.is_dir() and not d.name.startswith(("filtered", "_", "inputs"))
    )
    frames = []
    n_filtered = n_unfiltered = 0
    for ep in eps:
        if ep not in figured or ep in SCALAR_DIAGNOSTIC_DROP:
            continue
        filt_dir = inp / "filtered_v1" / ep
        use_filtered = filter_variant == "filtered_v1" and filt_dir.is_dir()
        src = filt_dir if use_filtered else inp / ep
        pqs = glob.glob(str(src / "*.parquet"))
        if not pqs:
            continue
        df = pd.concat([pd.read_parquet(p) for p in pqs], ignore_index=True)
        if "epoch_label" not in df.columns or "value" not in df.columns:
            continue
        sess = _persession_from_pertrial(df, ep)
        if not sess.empty:
            sess["family"] = "scalar"
            frames.append(sess)
            n_filtered += int(use_filtered)
            n_unfiltered += int(not use_filtered)
    log.info("scalar family: %d endpoints (%d filtered_v1, %d unfiltered) "
             "after figured-scope + diagnostic drop", n_filtered + n_unfiltered,
             n_filtered, n_unfiltered)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def load_phase_aligned_family(
    resolver: PathResolver, animals: list[str], *,
    trial_filter=None, kept_keys: set | None = None,
) -> pd.DataFrame:
    """First-lick phase-aligned endpoints (per-trial JSON) -> per-session.

    When ``kept_keys`` is provided the per-trial rows are gated by the same
    trial+session DLC filter the v3 figures apply (q13b), so the stats match
    the canonical filtered figures rather than all sessions.
    """
    import datetime

    from stroke_orofacial.cohort_enumeration import yaml_driven_session_dates
    from stroke_orofacial.stats.recovery_epochs import assign_epoch_label
    amap = load_animals()
    do_filter = kept_keys is not None
    rows = []
    for animal in animals:
        info = amap.get(animal)
        if info is None:
            continue
        for d in yaml_driven_session_dates(animal):
            if isinstance(d, str):
                d = datetime.date.fromisoformat(d)
            ep_long = load_session_phase_aligned_endpoints_long(animal, d, resolver)
            if ep_long.empty:
                continue
            if do_filter:
                licks = build_dlc_pertrial_lick_frame(animal, d, resolver)
                if licks is not None and "trial_index" in ep_long.columns:
                    ep_long = ep_long.merge(
                        licks[["animal", "date", "side", "trial_index",
                               "n_licks_per_trial"]],
                        on=["animal", "date", "side", "trial_index"], how="left")
                ep_long = _apply_trial_then_session_filter(
                    ep_long, trial_filter, kept_keys)
                if ep_long.empty:
                    continue
            # The JSON loader doesn't carry epoch_label; derive from rel_day.
            rel = int(info.day_relative_to_stroke(d))
            epoch = assign_epoch_label(rel, info)
            for ep in PHASE_ALIGNED_ENDPOINTS:
                sub = ep_long[ep_long["endpoint"] == ep]
                if sub.empty:
                    continue
                rows.append(_persession_from_pertrial(
                    sub.assign(is_sentinel=False, epoch_label=epoch), ep))
    if not rows:
        return pd.DataFrame()
    out = pd.concat(rows, ignore_index=True)
    out["family"] = "phase_aligned_endpoints"
    return out


def load_all_licks_endpoints_family(
    resolver: PathResolver, *, trial_filter=None, kept_keys: set | None = None,
) -> pd.DataFrame:
    """Precomputed per-LICK all-licks endpoints (stats/inputs/all_licks_endpoints/)
    -> per-session long-form. Requires precompute_all_licks_endpoints.py to have
    run; returns empty otherwise. ``kept_keys`` applies the session filter (and
    the trial filter where ``n_licks_per_trial`` is present) so the all-licks
    stats match the canonical v3 figures."""
    root = resolver.pipeline_outputs(subdir="stats") / "inputs" / "all_licks_endpoints"
    if not root.exists():
        return pd.DataFrame()
    frames = []
    for ep_dir in sorted(root.iterdir()):
        if not ep_dir.is_dir():
            continue
        pqs = glob.glob(str(ep_dir / "*.parquet"))
        if not pqs:
            continue
        df = pd.concat([pd.read_parquet(p) for p in pqs], ignore_index=True)
        if kept_keys is not None:
            df = _apply_trial_then_session_filter(df, trial_filter, kept_keys)
            if df.empty:
                continue
        sess = _persession_from_pertrial(df, ep_dir.name)
        if not sess.empty:
            sess["family"] = "phase_aligned_endpoints_all_licks"
            frames.append(sess)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def load_wasserstein_family(resolver: PathResolver) -> pd.DataFrame:
    """Persisted per-session W-2 (absolute + normalized) -> per-session long-form."""
    res = resolver.pipeline_outputs(subdir="stats") / "results" / "phenotype_pooled"
    frames = []
    for fam, sub in [("wasserstein_recovery", "wasserstein_recovery"),
                     ("normalized_wasserstein_recovery",
                      "normalized_wasserstein_recovery")]:
        for f in glob.glob(str(res / sub / "**" / "wasserstein_recovery_per_session.csv"),
                           recursive=True):
            df = pd.read_csv(f)
            if df.empty or "w2" not in df.columns:
                continue
            keep = df.rename(columns={"endpoint": "bodypart"})[
                ["animal", "date", "side", "bodypart", "epoch_label", "w2"]].copy()
            keep = keep.rename(columns={"w2": "value"})
            keep["endpoint"] = keep["bodypart"]
            keep["family"] = fam
            frames.append(keep.drop(columns=["bodypart"]))
    if not frames:
        return pd.DataFrame()
    out = pd.concat(frames, ignore_index=True)
    # De-duplicate scope-independent per-session rows. Each producer writes an
    # IDENTICAL per_session file under every phenotype-scope dir (``4_tier/``,
    # ``3_tier/``) because the write sits inside the per-scope emit loop, and
    # the recursive ``**`` glob above reads every copy. The W-2 per-session
    # value is scope-independent (phenotype tiering only affects downstream
    # pooling, not the raw session W-2), so without this dedup every session
    # would be counted once PER scope dir — inflating n and shrinking CIs/
    # p-values in the pooled stats. Key on the natural session identity; one
    # row per (animal, date, side, endpoint, epoch_label, family) is correct.
    key = ["animal", "date", "side", "endpoint", "epoch_label", "family"]
    return out.drop_duplicates(subset=key, ignore_index=True)


def _attach_phenotype(df: pd.DataFrame, scope: str) -> pd.DataFrame:
    amap = {a: i.recovery_phenotype for a, i in load_animals().items()}
    df = df.copy()
    df["recovery_phenotype"] = df["animal"].map(amap)
    df = df.dropna(subset=["recovery_phenotype"])
    if scope == "3_tier":
        df["recovery_phenotype"] = df["recovery_phenotype"].map(
            lambda p: PHENOTYPE_REMAP_3TIER.get(p, p))
    return df


def _run_family(long_df: pd.DataFrame, *, n_boot: int) -> pd.DataFrame:
    """Engine over every (family, endpoint, phenotype, side, post_epoch) cell."""
    out_rows: list[dict] = []
    # Group by (family, endpoint): the first-lick and all-licks phase-aligned
    # families share endpoint names (peak_angle_deg, ...), so grouping by
    # endpoint alone would collide them into one cell.
    for (family, endpoint), ep_df in long_df.groupby(["family", "endpoint"]):
        cross_cue = endpoint in CROSS_CUE_ENDPOINTS
        # Δ-from-pre per (animal, side) [or per animal if cross-cue].
        grp = ("animal",) if cross_cue else ("animal", "side")
        norm = SUITE.normalize_to_pre_baseline(ep_df, group_cols=grp, value_col="value")
        sides = ["both"] if cross_cue else sorted(
            s for s in norm["side"].dropna().unique())
        for side in sides:
            side_df = norm if cross_cue else norm[norm["side"] == side]
            for phen, phen_df in side_df.groupby("recovery_phenotype"):
                for post_epoch in SUITE.POST_EPOCHS:
                    if post_epoch not in set(phen_df["epoch_label"]):
                        continue
                    res = SUITE.compute_endpoint_cell_suite(
                        phen_df, post_epoch=post_epoch, value_col="value_norm",
                        n_boot=n_boot)
                    res.update({"family": family, "endpoint": endpoint,
                                "recovery_phenotype": phen, "side": side})
                    out_rows.append(res)
    return pd.DataFrame(out_rows)


def _attach_headline(res: pd.DataFrame) -> pd.DataFrame:
    """Add the per-family headline-test columns + a bootstrap-agreement flag.

    Session-MWU is the headline for the per-session scalar/endpoint families
    (LMM is non-estimable at 2-4 animals/phenotype); the LMM/Wilcoxon FDR
    columns remain as supporting evidence. ``headline_sig`` is the FDR-corrected
    headline call; ``headline_and_bootstrap_sig`` flags cells where the headline
    AND the session-bootstrap CI both reject the null.
    """
    res = res.copy()
    is_mwu = res["family"].isin(MWU_HEADLINE_FAMILIES)
    res["headline_test"] = np.where(is_mwu, "mwu_session", "lmm")
    res["headline_p_fdr"] = np.where(
        is_mwu, res.get("mwu_session_p_fdr"), res.get("lmm_p_fdr"))
    res["headline_sig"] = res["headline_p_fdr"] < 0.05
    res["headline_and_bootstrap_sig"] = (
        res["headline_sig"] & res.get("excludes_zero", False).astype(bool))
    return res


def _write_highlighted_xlsx(res: pd.DataFrame, out_path: Path) -> None:
    """Emit an .xlsx with per-test-column green(reject)/red(non-reject) fills.

    For each FDR test column (Wilcoxon / session-MWU / LMM): green if p<0.05
    (reject null), red if >=0.05, grey if NaN (e.g. LMM convergence_failed).
    The bootstrap ``excludes_zero`` bool: green if True, red if False. The
    headline columns get the same green/red treatment. (Priya 2026-06-01.)
    """
    from openpyxl.styles import PatternFill
    GREEN = PatternFill("solid", fgColor="C6EFCE")
    RED = PatternFill("solid", fgColor="FFC7CE")
    GREY = PatternFill("solid", fgColor="D9D9D9")
    cols = list(res.columns)
    pcols = [c for c in ("headline_p_fdr", "wilcoxon_p_fdr",
                         "mwu_session_p_fdr", "lmm_p_fdr") if c in cols]
    boolcols = [c for c in ("excludes_zero", "headline_sig",
                            "headline_and_bootstrap_sig") if c in cols]
    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        res.to_excel(xw, index=False, sheet_name="stats")
        ws = xw.sheets["stats"]
        ws.freeze_panes = "A2"
        for c in pcols:
            ci = cols.index(c) + 1
            for r in range(len(res)):
                v = res.iloc[r][c]
                cell = ws.cell(row=r + 2, column=ci)
                cell.fill = GREY if pd.isna(v) else (
                    GREEN if v < 0.05 else RED)
        for c in boolcols:
            ci = cols.index(c) + 1
            for r in range(len(res)):
                v = res.iloc[r][c]
                cell = ws.cell(row=r + 2, column=ci)
                cell.fill = GREEN if bool(v) else RED


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("--families", nargs="*",
                   default=["scalar", "phase_aligned_endpoints",
                            "phase_aligned_endpoints_all_licks", "wasserstein"])
    p.add_argument("--n-boot", type=int, default=5000)
    p.add_argument("--scopes", nargs="*", default=["4_tier", "3_tier"])
    p.add_argument("--filter-variant", choices=("filtered_v1", "none"),
                   default="filtered_v1",
                   help=("filtered_v1 (default): scalars read filtered_v1 "
                         "inputs where available + phase-aligned families apply "
                         "the trial+session DLC filter, matching the canonical "
                         "v3 figures. Spatial/wasserstein stays unfiltered. "
                         "none: all sessions (legacy)."))
    args = p.parse_args(argv)

    resolver = PathResolver()
    animals = sorted(load_animals().keys())
    filtered = args.filter_variant == "filtered_v1"

    # Build the DLC trial+session filter ONCE (shared by both phase-aligned
    # loaders). Unfiltered for --filter-variant none.
    trial_filter = kept_keys = None
    if filtered:
        log.info("building DLC trial+session filter (filtered_v1) ...")
        trial_filter, kept_keys = _build_dlc_filter(animals, resolver)

    base_long = []
    if "scalar" in args.families:
        log.info("loading scalar family ...")
        base_long.append(load_scalar_family(resolver, filter_variant=args.filter_variant))
    if "phase_aligned_endpoints" in args.families:
        log.info("loading phase_aligned family ...")
        base_long.append(load_phase_aligned_family(
            resolver, animals, trial_filter=trial_filter, kept_keys=kept_keys))
    if "phase_aligned_endpoints_all_licks" in args.families:
        log.info("loading all-licks endpoints family ...")
        base_long.append(load_all_licks_endpoints_family(
            resolver, trial_filter=trial_filter, kept_keys=kept_keys))
    if "wasserstein" in args.families:
        log.info("loading wasserstein family (always unfiltered — spatial) ...")
        base_long.append(load_wasserstein_family(resolver))
    base_long = [d for d in base_long if not d.empty]
    if not base_long:
        log.error("no family data loaded")
        return 1
    long_all = pd.concat(base_long, ignore_index=True)
    log.info("per-session rows: %d across %d endpoints (filter_variant=%s)",
             len(long_all), long_all["endpoint"].nunique(), args.filter_variant)

    for scope in args.scopes:
        scoped = _attach_phenotype(long_all, scope)
        res = _run_family(scoped, n_boot=args.n_boot)
        if res.empty:
            log.warning("[%s] no results", scope)
            continue
        # BH-FDR within family, per test column.
        for pcol, adj in [("wilcoxon_p", "wilcoxon_p_fdr"),
                          ("mwu_session_p", "mwu_session_p_fdr"),
                          ("lmm_p", "lmm_p_fdr")]:
            res = res.rename(columns={pcol: "p_value"})
            res = apply_correction_to_results_df(res, group_cols=("family",))
            res = res.rename(columns={"p_value": pcol, "p_adj": adj})
            res = res.drop(columns=["reject"], errors="ignore")
        res = _attach_headline(res)
        out_dir = (resolver.pipeline_outputs(subdir="stats") / "results"
                   / "phenotype_pooled" / scope)
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / "pooled_phenotype_stats.csv"
        res.to_csv(out_path, index=False)
        _write_highlighted_xlsx(res, out_dir / "pooled_phenotype_stats.xlsx")
        n_head = int(res["headline_sig"].sum())
        n_both = int(res["headline_and_bootstrap_sig"].sum())
        log.info("[%s] %d rows | %d headline-sig (MWU-FDR) | %d also bootstrap-sig "
                 "-> %s (+ .xlsx)", scope, len(res), n_head, n_both, out_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
